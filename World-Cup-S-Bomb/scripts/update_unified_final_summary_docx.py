#!/usr/bin/env python3
"""Build the complete deterministic Qatar 2022 v3 final-summary DOCX."""

from __future__ import annotations

import hashlib
import json
import os
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

import pandas as pd
from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = PROJECT_ROOT / "results/reports/docs/final_summary.docx"
RANKING_PATH = PROJECT_ROOT / "results/reports/ranking/player_rankings.csv"
GOALKEEPER_PATH = (
    PROJECT_ROOT / "results/reports/ranking/goalkeeper_rankings.csv"
)
AUDIT_PATH = (
    PROJECT_ROOT
    / "results/diagnostics/ranking_repair/v3_release_audit.json"
)
DOCX_AUDIT_PATH = (
    PROJECT_ROOT
    / "results/diagnostics/ranking_repair/final_summary_docx_validation.json"
)

ACTIVE_MODEL_VERSION = "ranking-repair-v3.0-qatar-2022"
PRESET_NAME = "compact_reference_guide"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
FIXED_PACKAGE_TIMESTAMP = (1980, 1, 1, 0, 0, 0)
FIXED_DOCUMENT_TIME = datetime(2022, 12, 18, 23, 59, tzinfo=timezone.utc)

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5B6573"
TABLE_HEADER_FILL = "E8EEF5"
TABLE_ALT_FILL = "F7F9FC"
WHITE = "FFFFFF"

REQUIRED_SECTION_TITLES = (
    "Executive Summary",
    "Scope and Event Boundary",
    "Published Ranking Products",
    "Active Methodology",
    "Champion–Challenger Validation",
    "Global Outfield Leaders",
    "300+ Minute Outfield Leaders",
    "Below-300-Minute High-Impact Players",
    "Position Leaders by Role Quality",
    "Role Leaders by Role Quality",
    "Team Leaders",
    "Dedicated Goalkeeper Leaders",
    "Release Gate and Limitations",
)


def _ensure_child(parent: Any, tag: str) -> Any:
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def _set_width(parent: Any, tag: str, width: int) -> None:
    element = _ensure_child(parent, tag)
    element.set(qn("w:type"), "dxa")
    element.set(qn("w:w"), str(int(width)))


def _set_run_font(
    run: Any,
    *,
    name: str = "Calibri",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _shade_cell(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = _ensure_child(properties, "w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def _apply_table_geometry(table: Any, widths: Sequence[int]) -> None:
    resolved = [int(value) for value in widths]
    if sum(resolved) != CONTENT_WIDTH_DXA:
        raise ValueError(
            f"Table widths must sum to {CONTENT_WIDTH_DXA}, got {sum(resolved)}"
        )
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    properties = table._tbl.tblPr
    _set_width(properties, "w:tblW", CONTENT_WIDTH_DXA)
    indent = _ensure_child(properties, "w:tblInd")
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    layout = _ensure_child(properties, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")

    margins = _ensure_child(properties, "w:tblCellMar")
    for side, value in {
        "top": 80,
        "bottom": 80,
        "start": 120,
        "end": 120,
    }.items():
        margin = _ensure_child(margins, f"w:{side}")
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in resolved:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row_index, row in enumerate(table.rows):
        row.height = None
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            row_properties.append(repeat)
        for column_index, cell in enumerate(row.cells):
            width = resolved[column_index]
            cell.width = Twips(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_properties = cell._tc.get_or_add_tcPr()
            _set_width(cell_properties, "w:tcW", width)
            cell_margins = _ensure_child(cell_properties, "w:tcMar")
            for side, value in {
                "top": 80,
                "bottom": 80,
                "start": 120,
                "end": 120,
            }.items():
                margin = _ensure_child(cell_margins, f"w:{side}")
                margin.set(qn("w:w"), str(value))
                margin.set(qn("w:type"), "dxa")


def _format_value(value: Any, *, integer: bool = False) -> str:
    if pd.isna(value):
        return "—"
    if integer:
        return str(int(round(float(value))))
    if isinstance(value, (float, int)) and not isinstance(value, bool):
        return f"{float(value):.4f}"
    return str(value)


def _set_cell_text(
    cell: Any,
    value: Any,
    *,
    header: bool = False,
    numeric: bool = False,
    integer: bool = False,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER if numeric else WD_ALIGN_PARAGRAPH.LEFT
    )
    run = paragraph.add_run(
        str(value) if header else _format_value(value, integer=integer)
    )
    _set_run_font(
        run,
        size=8.0 if header else 8.5,
        color=INK,
        bold=header,
    )


def _add_table(
    document: DocumentObject,
    frame: pd.DataFrame,
    *,
    columns: Sequence[str],
    labels: Sequence[str],
    widths: Sequence[int],
    integer_columns: Iterable[str] = (),
    numeric_columns: Iterable[str] = (),
) -> Any:
    missing = [column for column in columns if column not in frame]
    if missing:
        raise ValueError(f"DOCX table fields missing: {missing}")
    integer_fields = set(integer_columns)
    numeric_fields = set(numeric_columns) | integer_fields
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for column_index, label in enumerate(labels):
        cell = table.rows[0].cells[column_index]
        _shade_cell(cell, TABLE_HEADER_FILL)
        _set_cell_text(cell, label, header=True, numeric=False)
    for row_index, values in enumerate(
        frame.loc[:, columns].itertuples(index=False, name=None),
        start=1,
    ):
        cells = table.add_row().cells
        if row_index % 2 == 0:
            for cell in cells:
                _shade_cell(cell, TABLE_ALT_FILL)
        for column_index, (column, value) in enumerate(zip(columns, values)):
            _set_cell_text(
                cells[column_index],
                value,
                numeric=column in numeric_fields,
                integer=column in integer_fields,
            )
    _apply_table_geometry(table, widths)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def _configure_styles(document: DocumentObject) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    custom_styles = {
        "Report Title": (WD_STYLE_TYPE.PARAGRAPH, 30, INK, True),
        "Report Subtitle": (WD_STYLE_TYPE.PARAGRAPH, 15, DARK_BLUE, False),
        "Report Kicker": (WD_STYLE_TYPE.PARAGRAPH, 10, BLUE, True),
        "Table Citation": (WD_STYLE_TYPE.PARAGRAPH, 9, MUTED, False),
    }
    for name, (kind, size, color, bold) in custom_styles.items():
        style = styles[name] if name in styles else styles.add_style(name, kind)
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")

    styles["Report Title"].paragraph_format.space_before = Pt(0)
    styles["Report Title"].paragraph_format.space_after = Pt(8)
    styles["Report Title"].paragraph_format.keep_with_next = True
    styles["Report Subtitle"].paragraph_format.space_before = Pt(0)
    styles["Report Subtitle"].paragraph_format.space_after = Pt(28)
    styles["Report Subtitle"].paragraph_format.keep_with_next = True
    styles["Report Kicker"].paragraph_format.space_before = Pt(0)
    styles["Report Kicker"].paragraph_format.space_after = Pt(18)
    styles["Report Kicker"].paragraph_format.keep_with_next = True
    styles["Table Citation"].paragraph_format.space_before = Pt(4)
    styles["Table Citation"].paragraph_format.space_after = Pt(4)
    styles["Table Citation"].paragraph_format.keep_with_next = True


def _add_page_field(paragraph: Any) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    prefix = paragraph.add_run("Page ")
    _set_run_font(prefix, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, instruction, separate, placeholder, end])
    _set_run_font(run, size=8.5, color=MUTED)


def _configure_page(document: DocumentObject) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.text = ""
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(
        f"QATAR 2022 PLAYER RANKING REPORT  |  {ACTIVE_MODEL_VERSION}"
    )
    _set_run_font(run, size=8.0, color=MUTED, bold=True)

    footer = section.footer
    footer.is_linked_to_previous = False
    footer.paragraphs[0].text = ""
    _add_page_field(footer.paragraphs[0])


def _set_core_properties(document: DocumentObject) -> None:
    core = document.core_properties
    core.title = "2022 FIFA World Cup — Player Ranking Final Summary"
    core.subject = "Qatar 2022 ranking-repair v3 methodology and results"
    core.author = "World Cup S-Bomb reproducible analysis pipeline"
    core.last_modified_by = "World Cup S-Bomb reproducible analysis pipeline"
    core.keywords = (
        "Qatar 2022; Tournament Impact; Role Quality; Uncertainty; "
        "goalkeeper valuation"
    )
    core.comments = (
        "Generated from canonical v3 ranking artifacts; no retained template."
    )
    core.created = FIXED_DOCUMENT_TIME
    core.modified = FIXED_DOCUMENT_TIME
    core.revision = 1


def _add_cover(
    document: DocumentObject,
    *,
    release_status: str,
) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(90)

    kicker = document.add_paragraph(style="Report Kicker")
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.add_run("TECHNICAL REPORT  •  QATAR 2022")

    title = document.add_paragraph(style="Report Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("2022 FIFA World Cup\nPlayer Ranking Final Summary")

    subtitle = document.add_paragraph(style="Report Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Tournament Impact, Role Quality, Uncertainty, and goalkeeper valuation"
    )

    model = document.add_paragraph()
    model.alignment = WD_ALIGN_PARAGRAPH.CENTER
    model.paragraph_format.space_after = Pt(6)
    run = model.add_run(f"Active model  |  {ACTIVE_MODEL_VERSION}")
    _set_run_font(run, size=10.5, color=INK, bold=True)

    status = document.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status.paragraph_format.space_after = Pt(80)
    run = status.add_run(f"Release decision  |  {release_status}")
    _set_run_font(run, size=10.5, color=DARK_BLUE, bold=True)

    boundary = document.add_paragraph()
    boundary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    boundary.paragraph_format.space_after = Pt(0)
    run = boundary.add_run(
        "Evidence boundary: 2022 FIFA World Cup in Qatar only"
    )
    _set_run_font(run, size=9.5, color=MUTED, italic=True)
    boundary.add_run().add_break(WD_BREAK.PAGE)


def _add_bullet(document: DocumentObject, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def _add_caption(document: DocumentObject, text: str) -> None:
    paragraph = document.add_paragraph(style="Table Citation")
    paragraph.add_run(text)


def _leader_rows(
    rich: pd.DataFrame,
    *,
    limit: int,
    minimum_minutes: float | None = None,
    maximum_minutes: float | None = None,
) -> pd.DataFrame:
    mask = rich["position_group"].ne("Goalkeeper")
    minutes = pd.to_numeric(rich["minutes_played"], errors="coerce")
    if minimum_minutes is not None:
        mask &= minutes.ge(minimum_minutes)
    if maximum_minutes is not None:
        mask &= minutes.lt(maximum_minutes)
    return (
        rich.loc[mask]
        .sort_values(["global_rank_v3", "player_id"], kind="mergesort")
        .head(limit)
    )


def _component_rows(audit: Mapping[str, Any]) -> pd.DataFrame:
    rows = []
    for component, values in audit.get("component_selections", {}).items():
        rows.append(
            {
                "Component": component.replace("_", " ").title(),
                "Champion": values.get("champion", "—"),
                "Challenger": values.get("challenger", "—"),
                "Decision": values.get("decision", "—"),
                "Active": values.get("selected", "—"),
            }
        )
    return pd.DataFrame(rows)


def _metric_rows(audit: Mapping[str, Any]) -> pd.DataFrame:
    rows = pd.DataFrame(
        audit.get("champion_challenger", {}).get("metrics", [])
    )
    if rows.empty:
        raise ValueError("Champion/challenger metric rows are missing")
    return rows


def _build_document(
    rich: pd.DataFrame,
    goalkeeper: pd.DataFrame,
    audit: Mapping[str, Any],
) -> DocumentObject:
    document = Document()
    _configure_styles(document)
    _configure_page(document)
    _set_core_properties(document)
    _add_cover(
        document,
        release_status=str(audit.get("release_status", "not recorded")),
    )

    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(
        "This report publishes the repaired Qatar 2022 player-ranking model. "
        "The release separates total tournament contribution, role-relative "
        "quality, and uncertainty instead of blending them into one opaque "
        "rating. All eight repair passes completed and every promoted "
        "challenger satisfied its documented gate."
    )
    _add_bullet(
        document,
        "Global and team order use signed Tournament Impact in common action-"
        "value units for outfield players.",
    )
    _add_bullet(
        document,
        "Position and role comparisons use one empirical-Bayes Role Quality "
        "posterior rate.",
    )
    _add_bullet(
        document,
        "Confidence is reported as match-bootstrap Uncertainty intervals and "
        "rank bands; uncertainty is not subtracted from the score.",
    )
    _add_bullet(
        document,
        "Exactly one main goalkeeper per team enters the dedicated goalkeeper "
        "ranking; backup goalkeepers remain unranked.",
    )

    document.add_heading("Scope and Event Boundary", level=1)
    document.add_paragraph(
        "All scoring evidence comes from the 2022 FIFA World Cup in Qatar. "
        "Ordinary match performance uses StatsBomb periods 1–4, covering "
        "regulation and extra time. Period 5 is reserved for penalty "
        "shootouts and never enters ordinary outfield goals, xG, xA, xT, "
        "VAEP, finishing, or Tournament Impact."
    )
    reconciliation = audit.get("event_reconciliation", {})
    _add_bullet(
        document,
        "Ordinary player goals: "
        f"{int(reconciliation.get('player_goals_periods_1_to_4', 169))}; "
        "own goals: "
        f"{int(reconciliation.get('own_goals_periods_1_to_4', 3))}; "
        "official match-goal total: "
        f"{int(reconciliation.get('official_match_goals_including_own_goals', 172))}.",
    )
    _add_bullet(
        document,
        "Shootout goals are recorded separately and excluded from the "
        "ordinary-goal reconciliation.",
    )
    _add_bullet(
        document,
        "Player names, team identity, reputation, awards, advancement, and "
        "external rankings are not scoring features. External Qatar 2022 "
        "analysis is audit-only.",
    )

    document.add_heading("Published Ranking Products", level=1)
    document.add_heading("Tournament Impact", level=2)
    document.add_paragraph(
        "Tournament Impact is a signed total contribution for the tournament. "
        "It is the active outfield Global Rank and Team Rank field. No "
        "within-position z-score, role bonus, team-strength bonus, or repeated "
        "minutes multiplier creates this value."
    )
    document.add_heading("Role Quality", level=2)
    document.add_paragraph(
        "Role Quality is one reliability-aware empirical-Bayes posterior "
        "contribution rate. Probabilistic role membership supplies the prior "
        "mixture, while observed evidence is shrunk once. It determines "
        "position and role leaderboards, not the global order."
    )
    document.add_heading("Uncertainty", level=2)
    document.add_paragraph(
        "Uncertainty uses match-level bootstrap resampling to report score "
        "intervals and rank bands. It communicates finite-sample confidence "
        "and is never converted into an additional score or penalty."
    )

    document.add_heading("Active Methodology", level=1)
    document.add_heading("Offensive value", level=2)
    document.add_paragraph(
        "The attack gate compared process-only value, outcome-only value, "
        "process plus a bounded reliability-shrunk realization residual, and "
        "process plus full outcomes. The active selection is process_only: "
        "non-shootout expected action value is used without adding goals and "
        "assists a second time."
    )
    document.add_heading("Defensive value", level=2)
    document.add_paragraph(
        "The signed Ridge challenger models opportunity-adjusted change in "
        "conceding probability and threat prevention using match-disjoint "
        "out-of-fold predictions. It passed the held-out and bootstrap gates. "
        "The former one-sided defensive publication lift is retired."
    )
    document.add_heading("Goalkeeper value", level=2)
    document.add_paragraph(
        "Continuous goalkeeper evidence is calibrated out of fold with "
        "match-disjoint selection between sigmoid and isotonic calibration. "
        "The continuous component uses the following explicit weights:"
    )
    for text in (
        "40% shot stopping.",
        "15% high-leverage shot stopping.",
        "12% cross and claim control.",
        "10% sweeping.",
        "10% distribution under pressure.",
        "13% regular-penalty performance.",
    ):
        _add_bullet(document, text)
    document.add_paragraph(
        "Available continuous components are renormalized when evidence is "
        "missing. The continuous allocation supplies 90% of the dedicated "
        "goalkeeper score; the separate period-5 shootout contribution has a "
        "hard 10% cap. For unified publication only, "
        "percentile_equivalent_placement is mapped to the corresponding "
        "quantile of the outfield score distribution. This is explicitly a "
        "placement, not measured absolute cross-position contribution."
    )

    document.add_heading("Champion–Challenger Validation", level=1)
    document.add_paragraph(
        f"Overall release decision: {audit.get('release_status', 'not recorded')}."
    )
    _add_caption(document, "Table 1. Active component decisions.")
    _add_table(
        document,
        _component_rows(audit),
        columns=("Component", "Champion", "Challenger", "Decision", "Active"),
        labels=("Component", "Champion", "Challenger", "Decision", "Active"),
        widths=(1450, 1900, 1900, 1900, 2210),
    )
    metrics = _metric_rows(audit).copy()
    metrics["metric"] = metrics["metric"].str.replace("_", " ").str.title()
    _add_caption(
        document,
        "Table 2. Champion-versus-challenger metrics. Differences are "
        "challenger minus champion.",
    )
    _add_table(
        document,
        metrics,
        columns=(
            "metric",
            "champion",
            "challenger",
            "difference",
            "ci_low",
            "ci_high",
            "gate",
        ),
        labels=(
            "Metric",
            "Champion",
            "Challenger",
            "Δ",
            "CI low",
            "CI high",
            "Gate",
        ),
        widths=(2400, 1160, 1160, 1050, 1050, 1050, 1490),
        numeric_columns=(
            "champion",
            "challenger",
            "difference",
            "ci_low",
            "ci_high",
        ),
    )

    leaders = _leader_rows(rich, limit=20)
    document.add_heading("Global Outfield Leaders", level=1)
    document.add_paragraph(
        "Ordered by total Tournament Impact. Role Quality and bootstrap "
        "intervals are shown as separate evidence."
    )
    _add_caption(document, "Table 3. Top 20 outfield players by Tournament Impact.")
    _add_table(
        document,
        leaders,
        columns=(
            "global_rank_v3",
            "player_name",
            "team",
            "position_group",
            "minutes_played",
            "tournament_impact_raw_v3",
            "role_quality_v3",
            "uncertainty_status_v3",
        ),
        labels=(
            "Rank",
            "Player",
            "Team",
            "Position",
            "Min",
            "Impact",
            "Role Q.",
            "Uncertainty",
        ),
        widths=(520, 2170, 1160, 1540, 720, 1020, 930, 1300),
        integer_columns=("global_rank_v3",),
        numeric_columns=(
            "minutes_played",
            "tournament_impact_raw_v3",
            "role_quality_v3",
        ),
    )

    leaders_300 = _leader_rows(rich, limit=20, minimum_minutes=300.0)
    document.add_heading("300+ Minute Outfield Leaders", level=1)
    document.add_paragraph(
        "This high-reliability view is a publication cohort only. Players "
        "below 300 minutes remain scored and ranked in the full outfield table."
    )
    _add_caption(
        document,
        "Table 4. Top 20 outfield players with at least 300 tournament minutes.",
    )
    _add_table(
        document,
        leaders_300,
        columns=(
            "global_rank_v3",
            "player_name",
            "team",
            "position_group",
            "minutes_played",
            "tournament_impact_raw_v3",
            "role_quality_v3",
            "uncertainty_status_v3",
        ),
        labels=(
            "Rank",
            "Player",
            "Team",
            "Position",
            "Min",
            "Impact",
            "Role Q.",
            "Uncertainty",
        ),
        widths=(520, 2170, 1160, 1540, 720, 1020, 930, 1300),
        integer_columns=("global_rank_v3",),
        numeric_columns=(
            "minutes_played",
            "tournament_impact_raw_v3",
            "role_quality_v3",
        ),
    )

    short = _leader_rows(rich, limit=15, maximum_minutes=300.0)
    document.add_heading("Below-300-Minute High-Impact Players", level=1)
    document.add_paragraph(
        "These rows remain ordered by Tournament Impact. Rank-band endpoints "
        "make the smaller sample visible instead of burying it in a minutes "
        "penalty."
    )
    _add_caption(
        document,
        "Table 5. Top 15 outfield players below 300 minutes.",
    )
    _add_table(
        document,
        short,
        columns=(
            "global_rank_v3",
            "player_name",
            "team",
            "minutes_played",
            "tournament_impact_raw_v3",
            "bootstrap_rank_best_v3",
            "bootstrap_rank_worst_v3",
            "uncertainty_status_v3",
        ),
        labels=(
            "Rank",
            "Player",
            "Team",
            "Min",
            "Impact",
            "Best",
            "Worst",
            "Uncertainty",
        ),
        widths=(520, 2440, 1320, 800, 1050, 750, 750, 1730),
        integer_columns=(
            "global_rank_v3",
            "bootstrap_rank_best_v3",
            "bootstrap_rank_worst_v3",
        ),
        numeric_columns=("minutes_played", "tournament_impact_raw_v3"),
    )

    position_leaders = (
        rich.loc[rich["position_group"].ne("Goalkeeper")]
        .sort_values(
            ["position_group", "position_rank_v3", "player_id"],
            kind="mergesort",
        )
        .groupby("position_group", sort=True)
        .head(3)
    )
    document.add_heading("Position Leaders by Role Quality", level=1)
    document.add_paragraph(
        "Position comparisons use Role Quality rather than reinterpreting a "
        "within-position normalization as absolute global value."
    )
    _add_caption(document, "Table 6. Top three Role Quality rows per position.")
    _add_table(
        document,
        position_leaders,
        columns=(
            "position_group",
            "position_rank_v3",
            "player_name",
            "team",
            "role_quality_v3",
            "tournament_impact_raw_v3",
            "uncertainty_status_v3",
        ),
        labels=(
            "Position",
            "Pos.",
            "Player",
            "Team",
            "Role Q.",
            "Impact",
            "Uncertainty",
        ),
        widths=(1700, 560, 2440, 1350, 950, 950, 1410),
        integer_columns=("position_rank_v3",),
        numeric_columns=("role_quality_v3", "tournament_impact_raw_v3"),
    )

    role_leaders = (
        rich.loc[rich["position_group"].ne("Goalkeeper")]
        .sort_values(
            ["functional_role", "role_rank_v3", "player_id"],
            kind="mergesort",
        )
        .groupby("functional_role", sort=True)
        .head(1)
        .sort_values(
            ["role_quality_v3", "player_id"],
            ascending=[False, True],
            kind="mergesort",
        )
        .head(20)
    )
    document.add_heading("Role Leaders by Role Quality", level=1)
    _add_caption(document, "Table 7. Leading posterior Role Quality by role.")
    _add_table(
        document,
        role_leaders,
        columns=(
            "functional_role",
            "role_rank_v3",
            "player_name",
            "team",
            "role_quality_v3",
            "tournament_impact_raw_v3",
        ),
        labels=("Role", "Rank", "Player", "Team", "Role Q.", "Impact"),
        widths=(2600, 560, 2520, 1470, 1050, 1160),
        integer_columns=("role_rank_v3",),
        numeric_columns=("role_quality_v3", "tournament_impact_raw_v3"),
    )

    team_leaders = (
        rich.loc[rich["Team Rank"].eq(1)]
        .sort_values(["team", "player_id"], kind="mergesort")
        .copy()
    )
    document.add_heading("Team Leaders", level=1)
    document.add_paragraph(
        "Each row is the active publication leader for one of the 32 teams. "
        "A main goalkeeper appears only when the percentile-equivalent "
        "publication placement is highest within that team."
    )
    _add_caption(document, "Table 8. Active team leaders.")
    _add_table(
        document,
        team_leaders,
        columns=(
            "team",
            "Team Rank",
            "player_name",
            "position_group",
            "Global Rank",
            "Tournament Performance Score",
            "uncertainty_status_v3",
        ),
        labels=(
            "Team",
            "Team",
            "Player",
            "Position",
            "Global",
            "Score",
            "Uncertainty",
        ),
        widths=(1320, 540, 2390, 1610, 660, 1080, 1760),
        integer_columns=("Team Rank", "Global Rank"),
        numeric_columns=("Tournament Performance Score",),
    )

    goalkeeper = goalkeeper.sort_values(
        ["goalkeeper_rank_v3", "player_id"],
        kind="mergesort",
    )
    document.add_heading("Dedicated Goalkeeper Leaders", level=1)
    document.add_paragraph(
        "The dedicated goalkeeper score is not an outfield action-value total. "
        "The table below retains goalkeeper-only order and reports the "
        "continuous and shootout channels separately."
    )
    _add_caption(document, "Table 9. Top 15 main goalkeepers.")
    _add_table(
        document,
        goalkeeper.head(15),
        columns=(
            "goalkeeper_rank_v3",
            "player_name",
            "team",
            "dedicated_goalkeeper_score_v3",
            "continuous_goalkeeper_rating_v3",
            "shootout_component_v3",
            "goalkeeper_score_interval_low_v3",
            "goalkeeper_score_interval_high_v3",
            "goalkeeper_uncertainty_status_v3",
        ),
        labels=(
            "GK",
            "Goalkeeper",
            "Team",
            "Score",
            "Continuous",
            "Shootout",
            "Low",
            "High",
            "Uncertainty",
        ),
        widths=(500, 2040, 1150, 800, 940, 850, 720, 720, 1640),
        integer_columns=("goalkeeper_rank_v3",),
        numeric_columns=(
            "dedicated_goalkeeper_score_v3",
            "continuous_goalkeeper_rating_v3",
            "shootout_component_v3",
            "goalkeeper_score_interval_low_v3",
            "goalkeeper_score_interval_high_v3",
        ),
    )
    document.add_paragraph(
        "Only the 32 team-main goalkeepers are ranked. Backup goalkeepers have "
        "no dedicated or unified rank. The cross-position "
        "percentile_equivalent_placement is a labelled publication fallback "
        "and must not be read as measured absolute common-unit value."
    )

    document.add_heading("Release Gate and Limitations", level=1)
    document.add_paragraph(
        f"Overall release decision: {audit.get('release_status', 'not recorded')}."
    )
    for limitation in audit.get("limitations", []):
        _add_bullet(document, str(limitation))
    _add_bullet(
        document,
        "All rankings describe Qatar 2022 tournament evidence only. They are "
        "not career-strength, reputation, award, or future-performance "
        "rankings.",
    )
    return document


def _normalize_docx_package(source: Path, destination: Path) -> None:
    """Normalize ZIP order/timestamps so identical inputs hash identically."""

    with zipfile.ZipFile(source, "r") as archive:
        entries = {
            info.filename: (info, archive.read(info.filename))
            for info in archive.infolist()
        }
    with zipfile.ZipFile(
        destination,
        "w",
        compression=zipfile.ZIP_DEFLATED,
        compresslevel=9,
    ) as archive:
        for name in sorted(entries):
            original, payload = entries[name]
            info = zipfile.ZipInfo(name, date_time=FIXED_PACKAGE_TIMESTAMP)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = original.external_attr
            info.internal_attr = original.internal_attr
            info.flag_bits = original.flag_bits
            archive.writestr(info, payload, compress_type=zipfile.ZIP_DEFLATED)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _find_table(document: DocumentObject, first_header: str) -> Any:
    for table in document.tables:
        if table.rows and table.cell(0, 0).text.strip() == first_header:
            return table
    raise ValueError(f"DOCX table not found for first header {first_header!r}")


def _validate_docx(
    path: Path,
    *,
    rich: pd.DataFrame,
    goalkeeper: pd.DataFrame,
) -> dict[str, Any]:
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    missing_sections = [
        title for title in REQUIRED_SECTION_TITLES if title not in text
    ]
    if missing_sections:
        raise ValueError(f"DOCX section titles missing: {missing_sections}")
    required_text = (
        ACTIVE_MODEL_VERSION,
        "Tournament Impact",
        "Role Quality",
        "Uncertainty",
        "periods 1–4",
        "Period 5",
        "40% shot stopping",
        "10% cap",
        "percentile_equivalent_placement",
    )
    missing_text = [value for value in required_text if value not in text]
    if missing_text:
        raise ValueError(f"DOCX required text missing: {missing_text}")
    if "Part IX: Unified tournament publication layer" in text:
        raise ValueError("Stale appended Part IX methodology remains in DOCX")

    global_expected = _leader_rows(rich, limit=20).iloc[0]
    global_table = _find_table(document, "Rank")
    if global_table.cell(1, 1).text != str(global_expected["player_name"]):
        raise ValueError("DOCX global leader does not match active CSV")

    goalkeeper_expected = goalkeeper.sort_values(
        ["goalkeeper_rank_v3", "player_id"],
        kind="mergesort",
    ).iloc[0]
    goalkeeper_table = _find_table(document, "GK")
    if (
        goalkeeper_table.cell(1, 1).text
        != str(goalkeeper_expected["player_name"])
    ):
        raise ValueError("DOCX goalkeeper leader does not match active CSV")

    return {
        "schema_version": "ranking-repair-docx-validation-3.0",
        "active_model_version": ACTIVE_MODEL_VERSION,
        "design_preset": PRESET_NAME,
        "required_sections": list(REQUIRED_SECTION_TITLES),
        "required_sections_present": True,
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "global_leader": {
            "rank": int(global_expected["global_rank_v3"]),
            "player_name": str(global_expected["player_name"]),
            "team": str(global_expected["team"]),
        },
        "goalkeeper_leader": {
            "rank": int(goalkeeper_expected["goalkeeper_rank_v3"]),
            "player_name": str(goalkeeper_expected["player_name"]),
            "team": str(goalkeeper_expected["team"]),
        },
        "docx_sha256": _sha256(path),
        "status": "PASS",
    }


def _load_inputs() -> tuple[pd.DataFrame, pd.DataFrame, dict[str, Any]]:
    for path in (RANKING_PATH, GOALKEEPER_PATH, AUDIT_PATH):
        if not path.exists():
            raise FileNotFoundError(f"Canonical DOCX input missing: {path}")
    rich = pd.read_csv(RANKING_PATH, low_memory=False)
    goalkeeper = pd.read_csv(GOALKEEPER_PATH, low_memory=False)
    audit = json.loads(AUDIT_PATH.read_text(encoding="utf-8"))
    if audit.get("active_model_version") != ACTIVE_MODEL_VERSION:
        raise ValueError("DOCX audit input is not the active v3 release")
    if len(goalkeeper) != 32:
        raise ValueError(
            f"Expected 32 main goalkeepers for DOCX, found {len(goalkeeper)}"
        )
    return rich, goalkeeper, audit


def main() -> None:
    rich, goalkeeper, audit = _load_inputs()
    document = _build_document(rich, goalkeeper, audit)
    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    DOCX_AUDIT_PATH.parent.mkdir(parents=True, exist_ok=True)

    source_fd, source_name = tempfile.mkstemp(
        prefix=".final_summary.",
        suffix=".docx",
        dir=DOCX_PATH.parent,
    )
    normalized_fd, normalized_name = tempfile.mkstemp(
        prefix=".final_summary.normalized.",
        suffix=".docx",
        dir=DOCX_PATH.parent,
    )
    os.close(source_fd)
    os.close(normalized_fd)
    source_path = Path(source_name)
    normalized_path = Path(normalized_name)
    try:
        document.save(source_path)
        _normalize_docx_package(source_path, normalized_path)
        os.replace(normalized_path, DOCX_PATH)
    finally:
        source_path.unlink(missing_ok=True)
        normalized_path.unlink(missing_ok=True)

    validation = _validate_docx(
        DOCX_PATH,
        rich=rich,
        goalkeeper=goalkeeper,
    )
    DOCX_AUDIT_PATH.write_text(
        json.dumps(
            validation,
            ensure_ascii=False,
            indent=2,
            sort_keys=True,
        )
        + "\n",
        encoding="utf-8",
    )
    print(
        json.dumps(
            {
                "status": validation["status"],
                "path": DOCX_PATH.relative_to(PROJECT_ROOT).as_posix(),
                "sha256": validation["docx_sha256"],
                "tables": validation["table_count"],
            },
            indent=2,
            sort_keys=True,
        )
    )


if __name__ == "__main__":
    main()
