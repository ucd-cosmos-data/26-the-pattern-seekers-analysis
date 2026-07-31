#!/usr/bin/env python3
"""Build the promoted v5 final summary DOCX from its canonical Markdown."""

from __future__ import annotations

import re
import os
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "results" / "reports" / "final_summary.md"
OUTPUT = ROOT / "results" / "Summary" / "final_summary.docx"
CANONICAL_OUTPUT = (
    ROOT / "results" / "reports" / "docs" / "final_summary.docx"
)
QA_DIR = ROOT / "results" / "Summary" / "_docx_qa"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(90, 99, 110)
LIGHT_FILL = "E8EEF5"
TOTAL_DXA = 9360


def _set_font(run, *, size: float, bold: bool = False, color=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TOTAL_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row._tr.trPr.append(cant_split)
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in(
                "w:tcW"
            )
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _widths(headers: list[str]) -> list[int]:
    count = len(headers)
    if count == 7:
        return [540, 2020, 1120, 780, 1620, 1620, 1660]
    if count == 5:
        return [720, 2460, 1450, 1850, 2880]
    if count == 4:
        return [900, 2800, 1900, 3760]
    base = TOTAL_DXA // count
    values = [base] * count
    values[-1] += TOTAL_DXA - sum(values)
    return values


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    headers = rows[0]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    _repeat_header(table.rows[0])
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell._tc.get_or_add_tcPr().append(
            OxmlElement("w:shd")
        )
        cell._tc.tcPr[-1].set(qn("w:fill"), LIGHT_FILL)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.LEFT
            if index in {1, 2, 3}
            else WD_ALIGN_PARAGRAPH.CENTER
        )
        _set_font(paragraph.add_run(text), size=8, bold=True)
    for values in rows[1:]:
        cells = table.add_row().cells
        for index, text in enumerate(values):
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
                if index in {1, 2, 3}
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            _set_font(paragraph.add_run(text), size=7.8)
    _set_table_geometry(table, _widths(headers))
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def _configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size, before, after, color in (
        ("Heading 1", 16, 18, 10, BLUE),
        ("Heading 2", 13, 14, 7, BLUE),
        ("Heading 3", 12, 10, 5, DARK_BLUE),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_font(
        header.add_run("WORLD CUP S-BOMB  |  QATAR 2022"),
        size=8.5,
        bold=True,
        color=GRAY,
    )
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_font(footer.add_run("Page "), size=8.5, color=GRAY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def build() -> Path:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    _configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(72)
    title.paragraph_format.space_after = Pt(8)
    _set_font(
        title.add_run("World Cup S-Bomb"),
        size=30,
        bold=True,
        color=RGBColor(32, 55, 72),
    )
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    _set_font(
        subtitle.add_run("Qatar 2022 Final Model Summary"),
        size=15,
        color=DARK_BLUE,
    )
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(72)
    _set_font(
        kicker.add_run(
            "Outfield ranking-repair v3 + Consolidated Goalkeeper Value v5"
        ),
        size=10.5,
        bold=True,
        color=GRAY,
    )

    index = 1
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("## "):
            doc.add_paragraph(line[3:], style="Heading 1")
            index += 1
            continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
            index += 1
            continue
        if line.startswith("|"):
            table_rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                values = [
                    value.strip()
                    for value in lines[index].strip().strip("|").split("|")
                ]
                if not all(re.fullmatch(r":?-+:?", value) for value in values):
                    table_rows.append(values)
                index += 1
            if table_rows:
                _add_table(doc, table_rows)
            continue
        paragraph_lines = [line]
        index += 1
        while (
            index < len(lines)
            and lines[index].strip()
            and not lines[index].strip().startswith(("#", "|"))
        ):
            paragraph_lines.append(lines[index].strip())
            index += 1
        paragraph = doc.add_paragraph(" ".join(paragraph_lines))
        paragraph.paragraph_format.widow_control = True

    doc.core_properties.title = "World Cup S-Bomb — Qatar 2022 Final Summary"
    doc.core_properties.subject = (
        "Promoted single-metric goalkeeper v5 release"
    )
    doc.core_properties.author = "World Cup S-Bomb"
    from update_unified_final_summary_docx import (
        _normalize_docx_package,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    CANONICAL_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    source_fd, source_name = tempfile.mkstemp(
        prefix=".goalkeeper_v5_final_summary.",
        suffix=".docx",
        dir=OUTPUT.parent,
    )
    normalized_fd, normalized_name = tempfile.mkstemp(
        prefix=".goalkeeper_v5_final_summary.normalized.",
        suffix=".docx",
        dir=OUTPUT.parent,
    )
    os.close(source_fd)
    os.close(normalized_fd)
    source_path = Path(source_name)
    normalized_path = Path(normalized_name)
    try:
        doc.save(source_path)
        _normalize_docx_package(source_path, normalized_path)
        content = normalized_path.read_bytes()
        OUTPUT.write_bytes(content)
        CANONICAL_OUTPUT.write_bytes(content)
    finally:
        source_path.unlink(missing_ok=True)
        normalized_path.unlink(missing_ok=True)
    return OUTPUT


if __name__ == "__main__":
    print(build())
