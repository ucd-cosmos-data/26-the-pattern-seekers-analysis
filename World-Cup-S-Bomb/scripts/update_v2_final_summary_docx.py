#!/usr/bin/env python3
"""Update the retained 33-section Word methodology with canonical v2 results."""

from __future__ import annotations

import json
import os
import re
import tempfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = PROJECT_ROOT / "results/Summary/final_summary.docx"
MODEL_SUMMARY = PROJECT_ROOT / "results/reports/model_summary.json"
PROVENANCE = (
    PROJECT_ROOT / "data/processed/player_evaluation_provenance.json"
)
AUDIT_PATH = PROJECT_ROOT / "results/diagnostics/v2_docx_audit.json"


def _find(document: Document, text: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Document paragraph not found: {text!r}")


def _find_prefix(document: Document, prefix: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Document paragraph prefix not found: {prefix!r}")


def _set(paragraph: Paragraph, text: str) -> Paragraph:
    paragraph.text = text
    return paragraph


def _after(
    paragraph: Paragraph,
    text: str,
    *,
    style: str = "Normal",
) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    inserted = Paragraph(element, paragraph._parent)
    inserted.style = style
    inserted.add_run(text)
    return inserted


def _metric(
    payload: dict[str, Any],
    head: str,
    name: str,
) -> float:
    return float(
        payload["role_aware_elastic_net"]["heads"][head]["metrics"][name]
    )


def main() -> None:
    summary = json.loads(MODEL_SUMMARY.read_text(encoding="utf-8"))
    provenance = json.loads(PROVENANCE.read_text(encoding="utf-8"))
    document = Document(DOCX_PATH)
    numbered_sections_before = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name == "Heading 2"
        and re.match(r"^\d+\.", paragraph.text)
    ]
    if len(numbered_sections_before) != 33:
        raise RuntimeError(
            f"Expected 33 numbered sections, found "
            f"{len(numbered_sections_before)}"
        )

    eligibility_anchor = _find(
        document,
        "These broad groups are later used for outfield rating priors and role naming.",
    )
    inserted = _after(
        eligibility_anchor,
        "Outfield ratings are computed for every player with at least 45 "
        "tournament minutes; goalkeeper ratings require at least 90 minutes. "
        "Neither cohort is filtered at 300 minutes before feature construction, "
        "model training, position-prior estimation, or team aggregation.",
    )
    _after(
        inserted,
        "RankingStatus is Ranked (300+ min), Ranked (180-299 min), or "
        "Coverage only (<180 min). Goalkeepers additionally use "
        "GKRankingStatus with a 270-minute primary threshold.",
    )

    time_bullet = _find(document, "•\tTime.")
    _set(time_bullet, "•\tMatch minute.")
    inserted = _after(
        time_bullet,
        "•\tPre-action score differential (own goals minus opponent goals).",
        style="Compact Bullet",
    )
    inserted = _after(
        inserted,
        "•\tPre-tournament opponent strength from the 6 October 2022 FIFA ranking.",
        style="Compact Bullet",
    )
    inserted = _after(
        inserted,
        "•\tGame phase: group stage versus knockout.",
        style="Compact Bullet",
    )
    context_gate = provenance.get("contextual_vaep_gate", {})
    selected_context = context_gate.get(
        "selected_feature_set",
        "not recorded",
    )
    _after(
        inserted,
        "The four context variables form a challenger feature set. It is "
        "accepted only when match-disjoint development OOF performance meets "
        "the documented non-inferiority tolerances; the untouched test is "
        "opened once after that choice. This run selected the "
        f"{selected_context} feature set.",
    )

    defending_anchor = _find(document, "•\tDuel-win rate.")
    inserted = _after(
        defending_anchor,
        "•\txD90 percentile within the broad position group.",
        style="Compact Bullet",
    )
    _after(
        inserted,
        "The xD layer divides the pitch into 6 columns by 8 rows. Each "
        "interception, clearance, tackle/duel, block, pressure, recovery, or "
        "goalkeeper action receives BaseThreat(zone) multiplied by a fixed "
        "action-impact factor; the player total is converted to xD per 90.",
    )

    _set(
        _find_prefix(document, "Completeness_i="),
        "BaseCompleteness_i = mu_i · Balance_i",
    )
    completeness_end = _find(
        document,
        "This avoids awarding a perfect completeness score to someone who is uniformly average: the score requires both high top-three quality and balance.",
    )
    inserted = _after(
        completeness_end,
        "A minutes factor equals 0.50 below 90 minutes, 0.75 from 90-179, "
        "0.90 from 180-299, and 1.00 from 300 minutes.",
    )
    _after(
        inserted,
        "Completeness_i = 0.60 · BaseCompleteness_i + "
        "0.40 · MinutesFactor_i",
        style="Equation Block",
    )

    _set(
        _find(
            document,
            "The primary VAEP component takes the stronger channel:",
        ),
        "The primary VAEP component uses explicit functional-role weights "
        "whose offense and defense shares sum to one:",
    )
    _set(
        _find_prefix(document, "VAEPComponent_i="),
        "VAEPComponent_i = w_off(role_i) · OffVAEPScaled_i + "
        "w_def(role_i) · DefVAEPScaled_i",
    )
    _set(
        _find(
            document,
            "Thus an elite defensive contributor does not have to match an attacker’s raw offensive value to receive a strong VAEP component.",
        ),
        "Starting offense/defense weights are 0.85/0.15 for progressive "
        "wingers and forwards, 0.70/0.30 for attacking wingbacks and wide "
        "creators, 0.55/0.45 for hybrid creators, 0.35/0.65 for ball-winners, "
        "and 0.25/0.75 for center backs and defensive fullbacks.",
    )
    _set(
        _find_prefix(document, "RoleAdjustedValue_i="),
        "RoleAdjustedValue_i = w_off(role_i) · OffElasticNetScaled_i + "
        "w_def(role_i) · DefElasticNetScaled_i",
    )
    _set(
        _find_prefix(document, "“Role-adjusted” here refers"),
        "Role labels choose the symmetric channel weights but never award "
        "points, add player-specific bonuses, or override observed metrics.",
    )

    weights = summary["rating_weights"]
    formula = (
        "OutfieldRaw_i = "
        f"{weights['vaep_90']:.4f} · VAEPComponent_i + "
        f"{weights['vaep_per_touch']:.4f} · VAEPPerTouchPct_i + "
        f"{weights['xt_90']:.4f} · xT90Pct_i + "
        f"{weights['role_adjusted_value']:.4f} · RoleAdjustedValue_i + "
        f"{weights['completeness_score']:.4f} · Completeness_i + "
        f"{weights['off_ball_score']:.4f} · OffBall_i"
    )
    _set(_find_prefix(document, "OutfieldRaw_i="), formula)
    raw_intro = _find(document, "The raw rating is:")
    calibration = summary["composite_calibration"]
    _after(
        raw_intro,
        "A positive ElasticNet tests these six components against a transparent "
        "team-importance proxy with team-disjoint GroupKFold. Learned "
        "coefficients are normalized to sum to one. The challenger is retained "
        "only if its OOF Spearman correlation is non-inferior to the incumbent. "
        f"This run selected {'the incumbent fallback' if calibration['fallback_to_incumbent'] else 'the learned weights'} "
        f"({calibration['oof_spearman']:.4f} challenger versus "
        f"{calibration['incumbent_spearman']:.4f} incumbent).",
    )
    _set(
        _find_prefix(document, "MinutesReliability_i="),
        "MinutesReliability_i = Minutes_i / (Minutes_i + 450)",
    )
    _set(
        _find(
            document,
            "Only players with at least 300 tournament minutes enter the published rankings.",
        ),
        "Ratings and team aggregates include the full 45-minute eligible "
        "outfield cohort. The 300-minute cutoff only defines the primary "
        "high-reliability table; lower-minute players retain ratings and "
        "explicit RankingStatus labels.",
    )

    penalty_end = _find(
        document,
        "Goalkeepers with no observed penalties have this metric missing.",
    )
    inserted = _after(
        penalty_end,
        "32.7 High-leverage saves",
        style="Heading 3",
    )
    inserted = _after(
        inserted,
        "A high-leverage shot is a non-shootout shot on target with StatsBomb "
        "shot xG greater than 0.30.",
    )
    _after(
        inserted,
        "HighLeverageSavePct_i = Saves_i(xG > 0.30) / "
        "ShotsOnTargetFaced_i(xG > 0.30)",
        style="Equation Block",
    )

    _set(
        _find(document, "The six equally weighted components are:"),
        "The seven weighted components are:",
    )
    goalkeeper_items = [
        ("1.\tGoals prevented proxy per 90.", "1.\tGoals prevented proxy per 90: 22.5%."),
        ("2.\tSave rate.", "2.\tSave rate: 13.5%."),
        ("3.\tCross-stopping rate.", "3.\tCross-stopping rate: 13.5%."),
        ("4.\tSweeper actions per 90.", "4.\tSweeper actions per 90: 13.5%."),
        ("5.\tDistribution under pressure.", "5.\tDistribution under pressure: 13.5%."),
        ("6.\tShrunk penalty save rate.", "6.\tShrunk penalty save rate: 13.5%."),
    ]
    last_item: Paragraph | None = None
    for old, new in goalkeeper_items:
        last_item = _set(_find(document, old), new)
    if last_item is None:
        raise RuntimeError("Goalkeeper component list is empty")
    _after(
        last_item,
        "7.\tHigh-leverage save percentage: 10.0%.",
        style="Compact Number",
    )
    _set(
        _find_prefix(document, "GKRaw_i="),
        "GKRaw_i = sum_m available(w_m · Percentile_m,i) / "
        "sum_m available(w_m)",
    )
    _set(
        _find_prefix(document, "GKCoverage_i="),
        "GKCoverage_i = M_i / 7",
    )
    _set(
        _find_prefix(document, "GKReliability_i="),
        "GKReliability_i = GKCoverage_i · Minutes_i / (Minutes_i + 450)",
    )
    goalkeeper_impl = _find(
        document,
        "The implementation is in goalkeeper_valuation.py.",
    )
    _after(
        goalkeeper_impl,
        "GKFinal is computed for all goalkeepers with at least 90 minutes. "
        "GKRankingStatus is Ranked (270+ min), Ranked (180-269 min), or "
        "Coverage only (<180 min). Missing components reduce coverage and are "
        "never converted to zero.",
    )

    _set(
        _find_prefix(
            document,
            "The most important distinction is that many columns",
        ),
        "The most important distinction is that many final-CSV columns are "
        "explanatory or role-discovery variables. Only the six calibrated "
        "outfield components directly determine OutfieldRaw; xD affects the "
        "defensive role dimension. Seven weighted goalkeeper components "
        "determine GKRaw. Reliability shrinkage and missing-feature coverage "
        "then determine the published final scores.",
    )
    _set(
        _find_prefix(
            document,
            "The final hierarchy is best understood",
        ),
        "The final hierarchy is a tournament-sample evaluation. It preserves "
        "match-disjoint action modeling, explicitly tests calibration changes "
        "against a non-inferiority gate, and publishes coverage/reliability "
        "statuses. Limited 360 coverage, the event-derived post-shot proxy, "
        "within-tournament percentiles, and the weak defensive ElasticNet "
        "head still prevent interpretation as permanent player ability.",
    )

    vaep = provenance["vaep_oof_metrics"]
    test = provenance["vaep_final_test_metrics"]
    metric_replacements = {
        "•\tROC-AUC: 0.9490": f"•\tROC-AUC: {float(vaep['roc_auc']):.4f}",
        "•\tPR-AUC: 0.0848": f"•\tPR-AUC: {float(vaep['pr_auc']):.4f}",
        "•\tBrier: 0.001123": f"•\tBrier: {float(vaep['brier_score']):.6f}",
        "•\tROC-AUC: 0.9672": f"•\tROC-AUC: {float(test['roc_auc']):.4f}",
        "•\tPR-AUC: 0.1499": f"•\tPR-AUC: {float(test['pr_auc']):.4f}",
        "•\tBrier: 0.001494": f"•\tBrier: {float(test['brier_score']):.6f}",
    }
    for old, new in metric_replacements.items():
        _set(_find(document, old), new)

    if len(document.tables) < 7:
        raise RuntimeError("Expected the retained seven document tables")
    performance = document.tables[4]
    for row_index, head in ((1, "offense"), (2, "defense")):
        performance.cell(row_index, 1).text = (
            f"{_metric(summary, head, 'rmse'):.4f}"
        )
        performance.cell(row_index, 2).text = (
            f"{_metric(summary, head, 'mae'):.4f}"
        )
        performance.cell(row_index, 3).text = (
            f"{_metric(summary, head, 'correlation'):.4f}"
        )
    components = document.tables[5]
    component_rows = {
        row.cells[0].text.strip(): row
        for row in components.rows[1:]
    }
    component_rows["VAEP per 90"].cells[2].text = (
        "Role-weighted independently scaled OFF and DEF VAEP using the "
        f"development-gated {selected_context} feature set"
    )
    component_rows["Role-adjusted value"].cells[2].text = (
        "Role-weighted independently scaled offensive and defensive "
        "ElasticNet predictions"
    )
    component_rows["Completeness"].cells[2].text = (
        "60% top-three quality/balance plus 40% minutes factor"
    )
    table_weight_names = {
        "VAEP per 90": "vaep_90",
        "VAEP per touch": "vaep_per_touch",
        "xT per 90": "xt_90",
        "Role-adjusted value": "role_adjusted_value",
        "Completeness": "completeness_score",
        "Off-ball score": "off_ball_score",
    }
    for row_name, weight_name in table_weight_names.items():
        component_rows[row_name].cells[1].text = (
            f"{100 * float(weights[weight_name]):.1f}%"
        )
    influence = document.tables[6]
    for values in (
        (
            "Pre-action score/time/opponent/phase",
            (
                "Active"
                if selected_context == "contextual"
                else "OOF challenger rejected"
            ),
            "Indirect",
            "No",
            "No",
            "No",
        ),
        (
            "xD spatial disruption",
            "Defensive channel",
            "Defensive inputs",
            "Defending",
            "No",
            "No",
        ),
        (
            "High-leverage saves",
            "No",
            "No",
            "No",
            "No",
            "10% GKRaw",
        ),
    ):
        cells = influence.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            cell.text = value

    numbered_sections_after = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name == "Heading 2"
        and re.match(r"^\d+\.", paragraph.text)
    ]
    if numbered_sections_after != numbered_sections_before:
        raise RuntimeError("The retained 33 numbered sections changed")
    all_text = "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )
    stale_phrases = [
        "six equally weighted components",
        "Minutes_i+300",
        "Only players with at least 300 tournament minutes enter",
        "max(\nOffVAEPScaled_i",
    ]
    stale_found = [
        phrase for phrase in stale_phrases if phrase in all_text
    ]
    if stale_found:
        raise RuntimeError(f"Stale methodology remains: {stale_found}")

    descriptor, temporary_name = tempfile.mkstemp(
        dir=DOCX_PATH.parent,
        prefix=f".{DOCX_PATH.name}.",
        suffix=".tmp.docx",
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        document.save(temporary)
        os.replace(temporary, DOCX_PATH)
    finally:
        temporary.unlink(missing_ok=True)
    AUDIT_PATH.parent.mkdir(parents=True, exist_ok=True)
    AUDIT_PATH.write_text(
        json.dumps(
            {
                "document": str(DOCX_PATH),
                "numbered_sections": len(numbered_sections_after),
                "tables": len(document.tables),
                "inline_shapes": len(document.inline_shapes),
                "stale_phrases_found": stale_found,
                "rating_weights": weights,
                "calibration_fallback": calibration[
                    "fallback_to_incumbent"
                ],
            },
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    print(AUDIT_PATH.read_text(encoding="utf-8"))


if __name__ == "__main__":
    main()
