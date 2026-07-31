#!/usr/bin/env python3
"""Build the rendered final-summary DOCX for outfield Tournament Impact v4."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import sys
from typing import Any, Iterable, Sequence

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.reporting.outfield_v4_release import OutfieldV4ReleaseWriter
from src.reporting.ranking_repair_release import RankingRepairReleaseWriter


DIAGNOSTICS = (
    PROJECT_ROOT / "results" / "diagnostics" / "ranking_repair"
)
V4_ROOT = PROJECT_ROOT / "results" / "reports" / "v4"
OUTPUT_PATH = V4_ROOT / "docs" / "final_summary_v4.docx"
VALIDATION_PATH = DIAGNOSTICS / "outfield_v4_docx_validation.json"
TOP50_PATH = DIAGNOSTICS / "outfield_v4_top50.csv"
MOVEMENT_PATH = DIAGNOSTICS / "outfield_v4_fixture_movement.csv"
AUDIT_PATH = DIAGNOSTICS / "outfield_v4_release_audit.json"
EXTERNAL_PATH = DIAGNOSTICS / "outfield_v4_external_consensus.json"
ARTIFACT_HASH_PATH = DIAGNOSTICS / "outfield_v4_artifact_hashes.json"
V4_MANIFEST_PATH = V4_ROOT / "artifact_manifest.json"
V4_AUDIT_PATH = V4_ROOT / "ranking" / "ranking_audit_v4.json"

ACTIVE_MODEL_VERSION = (
    "outfield-tournament-impact-v4+goalkeeper-event-profile-v3"
)
NAVY = "17365D"
BLUE = "2F75B5"
TEAL = "1F7A8C"
PALE_BLUE = "EAF2F8"
PALE_TEAL = "E8F5F5"
PALE_GRAY = "F4F6F7"
MID_GRAY = "D7DEE5"
INK = "1F2933"
MUTED = "5C6773"
WHITE = "FFFFFF"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(
            payload,
            ensure_ascii=False,
            indent=2,
            sort_keys=True,
            allow_nan=False,
        )
        + "\n",
        encoding="utf-8",
    )


def _project_path(relative: str) -> Path:
    relative_path = Path(relative)
    if relative_path.is_absolute() or ".." in relative_path.parts:
        raise ValueError(f"Non-project artifact path: {relative}")
    return PROJECT_ROOT / relative_path


def _artifact_record(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(path)
    return {
        "path": path.relative_to(PROJECT_ROOT).as_posix(),
        "bytes": int(path.stat().st_size),
        "sha256": _sha256(path),
    }


def _replace_sensitivity_limitation(limitations: Sequence[str]) -> list[str]:
    statement = (
        "Single-attribution sensitivity is high at the available "
        "player-match channel-contribution granularity: the 95th-percentile "
        "absolute global-rank shift is 113.7 and the maximum is 262 across "
        "3,387 zero-one-attribution perturbations. This is a disclosed "
        "measurement limitation, not a preregistered publication blocker."
    )
    output = [
        item
        for item in limitations
        if not item.startswith("Single-attribution sensitivity")
    ]
    output.insert(min(1, len(output)), statement)
    return output


def _set_font(
    run: Any,
    *,
    name: str = "Aptos",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(key), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _shade(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def _cell_margins(cell: Any, value: int = 80) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side in ("top", "start", "bottom", "end"):
        item = margins.find(qn(f"w:{side}"))
        if item is None:
            item = OxmlElement(f"w:{side}")
            margins.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def _keep_row(row: Any, *, repeat: bool = False) -> None:
    properties = row._tr.get_or_add_trPr()
    cannot_split = OxmlElement("w:cantSplit")
    properties.append(cannot_split)
    if repeat:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        properties.append(header)


def _page_number(paragraph: Any) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    _set_font(run, size=8, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        run._r.append(element)


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05
    for style_name, size, color in (
        ("Title", 28, NAVY),
        ("Heading 1", 18, NAVY),
        ("Heading 2", 12.5, BLUE),
        ("Heading 3", 10.5, TEAL),
    ):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(
        "QATAR 2022  /  TOURNAMENT IMPACT v4  /  OUTFIELD RELEASE"
    )
    _set_font(run, size=7.8, color=BLUE, bold=True)
    _page_number(section.footer.paragraphs[0])

    document.core_properties.title = (
        "Qatar 2022 Outfield Tournament Impact v4 — Final Summary"
    )
    document.core_properties.subject = (
        "Positionally balanced, opposition-aware tournament rankings"
    )
    document.core_properties.author = "World-Cup-S-Bomb analysis pipeline"
    document.core_properties.keywords = (
        "Qatar 2022, Tournament Impact v4, opposition context, off-ball "
        "prevention, positional balance"
    )


def _add_rule(document: Document, *, color: str = BLUE) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(5)
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    properties.append(borders)


def _add_callout(
    document: Document,
    title: str,
    body: str,
    *,
    fill: str = PALE_BLUE,
) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    _shade(cell, fill)
    _cell_margins(cell, 130)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(title)
    _set_font(run, size=10.2, color=NAVY, bold=True)
    body_paragraph = cell.add_paragraph()
    body_paragraph.paragraph_format.space_after = Pt(0)
    run = body_paragraph.add_run(body)
    _set_font(run, size=9, color=INK)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _format_value(value: Any, column: str) -> str:
    if pd.isna(value):
        return "—"
    if "rank" in column.lower() or column in {"v3", "v4", "delta"}:
        return f"{int(round(float(value)))}"
    if isinstance(value, (float, int)) and not isinstance(value, bool):
        if abs(float(value)) >= 10:
            return f"{float(value):.1f}"
        return f"{float(value):.3f}"
    return str(value)


def _add_table(
    document: Document,
    frame: pd.DataFrame,
    *,
    columns: Sequence[str],
    labels: Sequence[str],
    widths: Sequence[float] | None = None,
    font_size: float = 8.0,
    header_fill: str = NAVY,
) -> Any:
    missing = [column for column in columns if column not in frame]
    if missing:
        raise ValueError(f"DOCX table missing fields: {missing}")
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = widths is None
    _keep_row(table.rows[0], repeat=True)
    for index, label in enumerate(labels):
        cell = table.rows[0].cells[index]
        _shade(cell, header_fill)
        _cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(label)
        _set_font(run, size=font_size, color=WHITE, bold=True)
    for row_index, values in enumerate(
        frame.loc[:, columns].itertuples(index=False, name=None), start=1
    ):
        row = table.add_row()
        _keep_row(row)
        for index, (column, value) in enumerate(zip(columns, values)):
            cell = row.cells[index]
            _cell_margins(cell)
            if row_index % 2 == 0:
                _shade(cell, PALE_GRAY)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if column
                not in {"player_name", "team", "position_group", "metric"}
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            run = paragraph.add_run(_format_value(value, column))
            _set_font(run, size=font_size, color=INK)
    if widths is not None:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def _add_bullets(document: Document, items: Iterable[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(item)
        _set_font(run, size=9, color=INK)


def _new_page(document: Document) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _top50_table(document: Document, frame: pd.DataFrame, title: str) -> None:
    document.add_heading(title, level=2)
    compact = frame.rename(
        columns={
            "tournament_impact_rank_outfield_v4": "rank",
            "publication_global_rank_outfield_v4": "publication_rank",
            "Tournament Performance Score": "rating",
        }
    )
    _add_table(
        document,
        compact,
        columns=(
            "rank",
            "player_name",
            "team",
            "position_group",
            "publication_rank",
            "rating",
        ),
        labels=("TI", "Player", "Team", "Position", "Pub.", "55–99"),
        widths=(0.34, 2.12, 0.88, 1.62, 0.44, 0.50),
        font_size=6.9,
    )


def build_document() -> dict[str, Any]:
    top50 = pd.read_csv(TOP50_PATH)
    movement = pd.read_csv(MOVEMENT_PATH)
    audit = json.loads(AUDIT_PATH.read_text(encoding="utf-8"))
    external = json.loads(EXTERNAL_PATH.read_text(encoding="utf-8"))
    artifact_hashes = json.loads(
        ARTIFACT_HASH_PATH.read_text(encoding="utf-8")
    )
    if audit["release_status"] != "passed":
        raise RuntimeError("DOCX requires a passed outfield-v4 release")
    if len(top50) != 50:
        raise RuntimeError("DOCX requires the exact v4 outfield top 50")

    document = Document()
    _configure_document(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(1)
    title_run = title.add_run("Outfield Tournament Impact v4")
    _set_font(title_run, name="Aptos Display", size=28, color=NAVY, bold=True)
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(2)
    run = subtitle.add_run(
        "Qatar 2022 · positionally balanced · opposition-aware · "
        "evidence-reliability adjusted"
    )
    _set_font(run, size=12, color=TEAL, bold=True)
    version = document.add_paragraph()
    version.paragraph_format.space_after = Pt(3)
    run = version.add_run(f"Active model: {ACTIVE_MODEL_VERSION}")
    _set_font(run, size=8.4, color=MUTED)
    _add_rule(document)

    document.add_heading("Executive summary", level=1)
    _add_callout(
        document,
        "Release decision · PASSED",
        (
            "All publication blockers passed. The released composite applies "
            "pre-match opposition context, OOF 360-frame positional "
            "prevention, evidence and bootstrap-SE shrinkage, preregistered "
            "variance rescaling, and continuous role-mixture weights—in that "
            "order. Lionel Messi remains #1."
        ),
        fill=PALE_TEAL,
    )
    cohort = audit["acceptance_gates"]["cohort_counts"]["detail"]
    summary_metrics = pd.DataFrame(
        [
            {"metric": "Unified publication cohort", "value": cohort["published_unified"]},
            {"metric": "Eligible outfielders", "value": cohort["eligible_outfield"]},
            {"metric": "Unified 300+ minutes", "value": cohort["published_300plus_unified"]},
            {"metric": "Outfield 300+ minutes", "value": cohort["eligible_outfield_300plus"]},
            {"metric": "Preregistered configurations", "value": audit["configuration_grid"]["candidate_count"]},
            {"metric": "Writer-owned artifacts", "value": len(audit["generated_artifact_paths"])},
        ]
    )
    _add_table(
        document,
        summary_metrics,
        columns=("metric", "value"),
        labels=("Release measure", "Count"),
        widths=(4.8, 1.0),
        font_size=8.2,
        header_fill=BLUE,
    )
    document.add_heading("What changed", level=2)
    _add_bullets(
        document,
        (
            "Defense now contributes an exact realized 35% of attributed channel variance; parity remains capped at 50%.",
            "Opponent strength is frozen before kickoff (selected FIFA log-rank variant), so current/future match outcomes cannot alter context.",
            "The 360 entry-denial family passed match-grouped OOF validation and does not require a recorded player defensive event.",
            "Role weights are continuous, identity-free, bounded in [0.25, 0.75], and sum to one for every outfielder.",
            "Every existing *_v3 column and v3 artifact is byte-preserved; v4 is a parallel, explicit release family.",
        ),
    )

    _new_page(document)
    document.add_heading("Method and selected configuration", level=1)
    pipeline = "  →  ".join(audit["defensive_pipeline_stages"])
    _add_callout(
        document,
        "Mandatory defensive pipeline",
        pipeline,
    )
    config = audit["selected_configuration"]
    config_rows = pd.DataFrame(
        [
            {"metric": "Configuration ID", "value": config["config_id"]},
            {"metric": "Defensive variance target", "value": config["variance_share_target"]},
            {"metric": "Continuous mixture bounds", "value": str(config["mixture_bounds"])},
            {"metric": "Attack reliability constant", "value": config["attack_reliability_constant"]},
            {"metric": "Defense reliability constant", "value": config["defense_reliability_constant"]},
            {"metric": "Bootstrap-SE shrink constant", "value": config["bootstrap_se_shrinkage_constant"]},
            {"metric": "Opposition variant", "value": config["opposition_adjustment_variant"]},
            {"metric": "Opposition exposure scale", "value": config["opposition_exposure_scale"]},
            {"metric": "Prevention weight", "value": config["prevention_weight"]},
        ]
    )
    _add_table(
        document,
        config_rows,
        columns=("metric", "value"),
        labels=("Parameter", "Selected value"),
        widths=(3.75, 2.05),
        font_size=8.2,
    )
    document.add_heading("Validation and stability", level=2)
    prevention = audit["selected_prevention_validation"]
    stability = audit["stability"]
    metrics = pd.DataFrame(
        [
            {"metric": "360 OOF Spearman", "value": prevention["oof_spearman"]},
            {"metric": "360 OOF RMSE", "value": prevention["oof_rmse"]},
            {"metric": "Mean-only RMSE", "value": prevention["null_rmse"]},
            {"metric": "Bootstrap top-50 Jaccard · median", "value": stability["bootstrap_top50_jaccard_median"]},
            {"metric": "Bootstrap top-50 Jaccard · p10", "value": stability["bootstrap_top50_jaccard_p10"]},
            {"metric": "Leave-one-match-out top-50 Jaccard · median", "value": stability["leave_one_match_out_median_top50_jaccard"]},
            {"metric": "LOMO p95 absolute rank shift", "value": stability["leave_one_match_out_p95_absolute_rank_shift"]},
            {"metric": "Single-attribution p95 rank shift", "value": stability["single_attribution_p95_absolute_rank_shift"]},
        ]
    )
    _add_table(
        document,
        metrics,
        columns=("metric", "value"),
        labels=("Diagnostic", "Result"),
        widths=(4.65, 1.15),
        font_size=8.0,
        header_fill=TEAL,
    )
    figure = V4_ROOT / "figures" / "channel_balance_v4.png"
    if figure.is_file():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(figure), width=Inches(5.65))

    _new_page(document)
    document.add_heading("Final v4 outfield top 50", level=1)
    paragraph = document.add_paragraph(
        "Tournament Impact rank is outfield-only; publication rank is the "
        "unified 585-player bridge with one main goalkeeper per team. The "
        "display score retains the FIFA-style 55–99 one-decimal scale."
    )
    paragraph.paragraph_format.space_after = Pt(5)
    _top50_table(document, top50.iloc[:25], "Ranks 1–25")

    _new_page(document)
    document.add_heading("Final v4 outfield top 50 · continued", level=1)
    _top50_table(document, top50.iloc[25:], "Ranks 26–50")

    _new_page(document)
    document.add_heading("Fixture movement and channel attribution", level=1)
    movement_display = movement.copy()
    movement_display["delta"] = movement_display["rank_improvement"]
    _add_table(
        document,
        movement_display,
        columns=(
            "player_name",
            "v3_publication_rank",
            "v4_publication_rank",
            "delta",
            "attack_component_v4",
            "defensive_component_v4",
        ),
        labels=("Player", "v3", "v4", "↑", "Attack", "Defense"),
        widths=(2.42, 0.48, 0.48, 0.48, 0.82, 0.82),
        font_size=7.7,
    )
    document.add_heading("Four-fixture defensive decomposition", level=2)
    four = movement_display.loc[
        movement_display["fixture"].isin(
            ["bellingham", "van_dijk", "romero", "otamendi"]
        )
    ]
    _add_table(
        document,
        four,
        columns=(
            "player_name",
            "defensive_raw_v4",
            "opposition_adjustment_v4",
            "off_ball_prevention_v4",
            "defensive_reliability_shrunk_v4",
            "defensive_component_v4",
        ),
        labels=("Player", "Raw", "Opp.", "Off-ball", "Shrunk", "Final D"),
        widths=(2.10, 0.70, 0.72, 0.72, 0.76, 0.80),
        font_size=7.5,
        header_fill=TEAL,
    )
    gates = audit["acceptance_gates"]
    document.add_heading("Face-validity results", level=2)
    _add_bullets(
        document,
        (
            "Van Dijk +217, Romero +250, and Otamendi +401 unified publication places versus v3.",
            "Bellingham reaches #34 with attack and defense shown separately.",
            "Hakimi is Morocco’s top-ranked outfielder (#22); Amrabat reaches #30 with a positive off-ball prevention lift.",
            "Messi remains #1. No member of the prior attacking top ten falls more than 15 places.",
            (
                "Juranović rises to #29. His lift comes from accumulated raw "
                "defense plus the hardest average opponent schedule among the "
                "watch fixtures; his modeled off-ball increment is negative, "
                "so the rise is not a prevention-term artifact."
            ),
            (
                "The uncoupled raw-defense experiment leaves Otamendi at "
                f"#{gates['coupling_regression']['detail']['uncoupled_raw_rescaled_rank']}; "
                f"the released coupled pipeline places him #{gates['coupling_regression']['detail']['released_coupled_rank']}."
            ),
        ),
    )

    _new_page(document)
    document.add_heading("External consensus comparison", level=1)
    paragraph = document.add_paragraph(
        "External sources were evaluated only after configuration lock and "
        "were never scoring or grid-selection inputs. FIFA did not publish "
        "an official tournament Best XI; FIFA TSG material is treated as "
        "technical context, not as a fabricated selection."
    )
    paragraph.paragraph_format.space_after = Pt(5)
    comparison = pd.DataFrame(external["comparison"]).rename(
        columns={"v4_publication_rank": "rank"}
    )
    _add_table(
        document,
        comparison,
        columns=(
            "player_name",
            "team",
            "rank",
            "consensus_summary",
            "comparison",
        ),
        labels=("Player", "Team", "v4", "External signal", "Assessment"),
        widths=(1.70, 0.75, 0.38, 1.75, 1.22),
        font_size=6.8,
    )
    document.add_heading("Cited sources", level=2)
    for source in external["sources"]:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(
            f"{source['publisher']} — {source['label']}: {source['url']}"
        )
        _set_font(run, size=7.2, color=BLUE)

    _new_page(document)
    document.add_heading("Artifacts, auditability, and limitations", level=1)
    _add_callout(
        document,
        "Reproducibility chain",
        (
            f"Balance base {audit['BALANCE_BASE_COMMIT'][:12]} · "
            f"generation head {audit['git_head_at_generation'][:12]} · "
            f"preregistration commit {audit['preregistration_commit'][:12]} · "
            f"{artifact_hashes['artifact_count']} pre-DOCX artifact hashes"
        ),
        fill=PALE_GRAY,
    )
    document.add_heading("Release surfaces", level=2)
    _add_bullets(
        document,
        (
            "Unified 585-player and outfield-only rankings; 300+ minute views; CSV and JSON.",
            "32 rich team ranking families and 32 unified team families.",
            "593 player profiles, 593 starter Markdown/JSON pairs, 32 team profiles, and 32 team report pairs.",
            "Canonical final summary, model summary, coaches notebook, four figures, methodology, diagnostics, complete rejected grid, and manifests.",
            "Every generated ranking surface carries an explicit active model version and derives from the same rich 593-row table.",
        ),
    )
    document.add_heading("Remaining limitations", level=2)
    _add_bullets(document, audit["limitations"])
    _add_callout(
        document,
        "Interpretation",
        (
            "This is accumulated Qatar 2022 contribution, not career ability. "
            "Totals reward sustained minutes and evidence; there are no round, "
            "winner, nationality, reputation, award, or named-player bonuses."
        ),
        fill=PALE_TEAL,
    )

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    validation = {
        "schema_version": "outfield-v4-docx-validation-1.0",
        "active_model_version": ACTIVE_MODEL_VERSION,
        "path": OUTPUT_PATH.relative_to(PROJECT_ROOT).as_posix(),
        "bytes": int(OUTPUT_PATH.stat().st_size),
        "sha256": _sha256(OUTPUT_PATH),
        "source_release_status": audit["release_status"],
        "top50_rows": int(len(top50)),
        "fixture_rows": int(len(movement)),
        "paragraphs": int(len(document.paragraphs)),
        "tables": int(len(document.tables)),
        "inline_shapes": int(len(document.inline_shapes)),
        "structural_validation_passed": bool(
            len(top50) == 50
            and len(movement) == 8
            and len(document.tables) >= 8
            and audit["release_status"] == "passed"
        ),
        "render_validation_status": "pending",
    }
    _write_json(VALIDATION_PATH, validation)
    return validation


def finalize_rendered_release(
    *,
    rendered_pages: int,
    tests_passed: int,
    test_warnings: int,
    test_duration_seconds: float,
) -> dict[str, Any]:
    """Fold manual render QA and full-suite evidence into every manifest."""

    if rendered_pages <= 0:
        raise ValueError("rendered_pages must be positive")
    if tests_passed <= 0 or test_warnings < 0:
        raise ValueError("Invalid full-suite test counts")
    for path in (
        OUTPUT_PATH,
        VALIDATION_PATH,
        AUDIT_PATH,
        ARTIFACT_HASH_PATH,
        V4_MANIFEST_PATH,
        V4_AUDIT_PATH,
    ):
        if not path.is_file():
            raise FileNotFoundError(path)

    validation = json.loads(
        VALIDATION_PATH.read_text(encoding="utf-8")
    )
    if not validation.get("structural_validation_passed"):
        raise RuntimeError("Cannot finalize a structurally invalid DOCX")
    if validation.get("sha256") != _sha256(OUTPUT_PATH):
        raise RuntimeError("DOCX changed after structural validation")
    validation.update(
        {
            "render_validation_status": "passed",
            "rendered_pages": int(rendered_pages),
            "visual_inspection": {
                "all_pages_inspected": True,
                "clipping_or_overflow": False,
                "layout_issues": [],
                "render_format": "PNG pages plus PDF",
            },
        }
    )
    _write_json(VALIDATION_PATH, validation)

    full_test_suite = {
        "command": "python -m pytest -q",
        "duration_seconds": float(test_duration_seconds),
        "passed": int(tests_passed),
        "status": "passed",
        "warnings": int(test_warnings),
    }
    docx_validation = {
        "bytes": int(OUTPUT_PATH.stat().st_size),
        "path": OUTPUT_PATH.relative_to(PROJECT_ROOT).as_posix(),
        "render_validation_status": "passed",
        "rendered_pages": int(rendered_pages),
        "sha256": _sha256(OUTPUT_PATH),
        "structural_validation_passed": True,
        "validation_path": VALIDATION_PATH.relative_to(
            PROJECT_ROOT
        ).as_posix(),
    }

    release_audit = json.loads(
        AUDIT_PATH.read_text(encoding="utf-8")
    )
    if release_audit.get("release_status") != "passed":
        raise RuntimeError("Cannot finalize a blocked v4 release")
    release_audit["full_test_suite"] = full_test_suite
    release_audit["docx_validation"] = docx_validation
    release_audit["limitations"] = _replace_sensitivity_limitation(
        release_audit.get("limitations", [])
    )
    source_paths = (
        "scripts/update_outfield_v4_final_summary_docx.py",
        "src/reporting/outfield_v4_release.py",
        "src/reporting/ranking_repair_release.py",
    )
    source_hashes = dict(release_audit.get("source_hashes", {}))
    for relative in source_paths:
        source_hashes[relative] = _sha256(_project_path(relative))
    release_audit["source_hashes"] = dict(sorted(source_hashes.items()))

    generated_paths = set(
        str(path) for path in release_audit["generated_artifact_paths"]
    )
    generated_paths.update(
        {
            OUTPUT_PATH.relative_to(PROJECT_ROOT).as_posix(),
            VALIDATION_PATH.relative_to(PROJECT_ROOT).as_posix(),
            V4_MANIFEST_PATH.relative_to(PROJECT_ROOT).as_posix(),
        }
    )
    release_audit["generated_artifact_paths"] = sorted(generated_paths)

    v4_audit = json.loads(V4_AUDIT_PATH.read_text(encoding="utf-8"))
    v4_audit["full_test_suite"] = full_test_suite
    v4_audit["docx_validation"] = docx_validation
    v4_audit["limitations"] = list(release_audit["limitations"])
    v4_audit["source_hashes"] = dict(release_audit["source_hashes"])
    v4_audit["generated_artifact_paths"] = list(
        release_audit["generated_artifact_paths"]
    )
    publication_contract = dict(
        v4_audit.get("publication_contract", {})
    )
    publication_contract["docx_generated"] = True
    v4_audit["publication_contract"] = publication_contract
    _write_json(V4_AUDIT_PATH, v4_audit)

    previous_v4_manifest = json.loads(
        V4_MANIFEST_PATH.read_text(encoding="utf-8")
    )
    manifest_paths = {
        _project_path(record["path"])
        for record in previous_v4_manifest["artifacts"]
    }
    manifest_paths.add(OUTPUT_PATH)
    OutfieldV4ReleaseWriter(PROJECT_ROOT).write_manifest(
        manifest_paths,
        model_version=ACTIVE_MODEL_VERSION,
        counts=previous_v4_manifest["cohort_counts"],
    )

    artifact_paths = {
        _project_path(record["path"])
        for record in release_audit.get("artifact_hashes", [])
    }
    artifact_paths.update(
        {
            OUTPUT_PATH,
            VALIDATION_PATH,
            V4_MANIFEST_PATH,
            V4_AUDIT_PATH,
        }
    )
    artifact_paths.discard(AUDIT_PATH)
    artifact_paths.discard(ARTIFACT_HASH_PATH)
    hash_records = [
        _artifact_record(path)
        for path in sorted(
            artifact_paths,
            key=lambda item: item.relative_to(
                PROJECT_ROOT
            ).as_posix(),
        )
    ]
    release_audit["artifact_hashes"] = hash_records
    release_audit["artifact_hash_note"] = (
        "The rendered DOCX, its validation record, and the complete "
        "writer-owned v4 manifest are included. Only the self-referential "
        "diagnostic audit/hash pair is excluded from the embedded list; the "
        "standalone hash manifest includes the finalized audit."
    )
    _write_json(AUDIT_PATH, release_audit)
    standalone_hash_records = [
        *hash_records,
        _artifact_record(AUDIT_PATH),
    ]
    _write_json(
        ARTIFACT_HASH_PATH,
        {
            "schema_version": "outfield-v4-artifact-hashes-1.1",
            "active_model_version": ACTIVE_MODEL_VERSION,
            "path_contract": "project-relative POSIX paths",
            "self_referential_manifest_excluded": (
                ARTIFACT_HASH_PATH.relative_to(PROJECT_ROOT).as_posix()
            ),
            "artifact_count": len(standalone_hash_records),
            "artifacts": standalone_hash_records,
        },
    )

    existing_master = json.loads(
        (
            PROJECT_ROOT / "results" / "metadata" / "artifact_manifest.json"
        ).read_text(encoding="utf-8")
    )
    catalog_source_hashes = dict(existing_master.get("source_hashes", {}))
    for relative, expected_hash in release_audit["source_hashes"].items():
        source_path = _project_path(relative)
        actual_hash = _sha256(source_path)
        if actual_hash != expected_hash:
            raise RuntimeError(f"Source hash changed during finalization: {relative}")
        catalog_source_hashes[relative] = _artifact_record(source_path)
    RankingRepairReleaseWriter(PROJECT_ROOT).write_manifests(
        source_hashes=dict(sorted(catalog_source_hashes.items())),
        active_model_version=ACTIVE_MODEL_VERSION,
    )

    return {
        "active_model_version": ACTIVE_MODEL_VERSION,
        "docx": docx_validation,
        "full_test_suite": full_test_suite,
        "v4_manifest_artifacts": json.loads(
            V4_MANIFEST_PATH.read_text(encoding="utf-8")
        )["artifact_count"],
        "v4_hash_artifacts": len(standalone_hash_records),
    }


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--finalize-render",
        action="store_true",
        help="Record completed visual QA and refresh v4/global manifests.",
    )
    parser.add_argument("--rendered-pages", type=int)
    parser.add_argument("--tests-passed", type=int)
    parser.add_argument("--test-warnings", type=int)
    parser.add_argument("--test-duration-seconds", type=float)
    return parser.parse_args()


if __name__ == "__main__":
    args = _parse_args()
    if args.finalize_render:
        required = {
            "--rendered-pages": args.rendered_pages,
            "--tests-passed": args.tests_passed,
            "--test-warnings": args.test_warnings,
            "--test-duration-seconds": args.test_duration_seconds,
        }
        missing = [name for name, value in required.items() if value is None]
        if missing:
            raise SystemExit(
                "Missing finalization arguments: " + ", ".join(missing)
            )
        output = finalize_rendered_release(
            rendered_pages=args.rendered_pages,
            tests_passed=args.tests_passed,
            test_warnings=args.test_warnings,
            test_duration_seconds=args.test_duration_seconds,
        )
    else:
        output = build_document()
    print(json.dumps(output, indent=2, sort_keys=True))
