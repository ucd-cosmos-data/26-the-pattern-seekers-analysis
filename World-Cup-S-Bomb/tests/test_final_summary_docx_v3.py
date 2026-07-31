"""Contract tests for the deterministic promoted goalkeeper v5 DOCX."""

from __future__ import annotations

import hashlib
import subprocess
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.oxml.ns import qn


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = PROJECT_ROOT / "results/reports/docs/final_summary.docx"
RANKING_PATH = PROJECT_ROOT / "results/reports/ranking/player_rankings.csv"
GOALKEEPER_PATH = (
    PROJECT_ROOT / "results/reports/ranking/goalkeeper_rankings.csv"
)
GENERATOR = PROJECT_ROOT / "scripts/build_goalkeeper_v5_final_docx.py"
ACTIVE_MODEL = "ranking-repair-v3.0-qatar-2022"
REQUIRED_SECTIONS = (
    "Active model",
    "Global outfield top 20",
    "Outfield players with 300+ minutes",
    "Goalkeeper ranking",
    "Goalkeeper methodology and validation",
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _table_by_header(document, header: str):
    return next(
        table
        for table in document.tables
        if table.rows and table.cell(0, 0).text.strip() == header
    )


def test_docx_is_complete_and_matches_active_rankings() -> None:
    document = Document(DOCX_PATH)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for title in REQUIRED_SECTIONS:
        assert title in text
    for phrase in (
        ACTIVE_MODEL,
        "Tournament Impact",
        "Role Quality",
        "Uncertainty",
        "periods 1–4",
        "Period 5",
        "Consolidated Goalkeeper Value v5",
        "one active metric",
        "All twelve hard promotion gates passed",
    ):
        assert phrase in text
    assert len(document.tables) == 3

    rankings = pd.read_csv(RANKING_PATH, low_memory=False)
    expected_global = (
        rankings.loc[rankings["publication_global_rank_v5"].notna()]
        .sort_values(
            ["publication_global_rank_v5", "player_id"],
            kind="mergesort",
        )
        .iloc[0]
    )
    global_table = _table_by_header(document, "Global Rank")
    assert global_table.cell(1, 0).text == str(
        int(expected_global["publication_global_rank_v5"])
    )
    assert global_table.cell(1, 1).text == expected_global["player_name"]
    assert global_table.cell(1, 2).text == expected_global["team"]

    goalkeepers = pd.read_csv(GOALKEEPER_PATH, low_memory=False).sort_values(
        ["goalkeeper_consolidated_value_rank_v5", "player_id"],
        kind="mergesort",
    )
    assert len(goalkeepers) == 32
    expected_goalkeeper = goalkeepers.iloc[0]
    goalkeeper_table = _table_by_header(document, "GK Rank")
    assert goalkeeper_table.cell(1, 0).text == str(
        int(expected_goalkeeper["goalkeeper_consolidated_value_rank_v5"])
    )
    assert goalkeeper_table.cell(1, 1).text == expected_goalkeeper[
        "player_name"
    ]
    assert goalkeeper_table.cell(1, 2).text == expected_goalkeeper["team"]


def test_docx_uses_compact_reference_geometry() -> None:
    document = Document(DOCX_PATH)
    section = document.sections[0]
    assert round(section.page_width.inches, 3) == 8.5
    assert round(section.page_height.inches, 3) == 11.0
    assert round(section.top_margin.inches, 3) == 1.0
    assert round(section.right_margin.inches, 3) == 1.0
    assert round(section.bottom_margin.inches, 3) == 1.0
    assert round(section.left_margin.inches, 3) == 1.0
    assert round(section.header_distance.inches, 3) == 0.492
    assert round(section.footer_distance.inches, 3) == 0.492

    normal = document.styles["Normal"]
    assert normal.font.name == "Calibri"
    assert normal.font.size.pt == 11
    assert normal.paragraph_format.space_after.pt == 6
    assert normal.paragraph_format.line_spacing == 1.25
    for style_name, size, before, after in (
        ("Heading 1", 16, 18, 10),
        ("Heading 2", 13, 14, 7),
        ("Heading 3", 12, 10, 5),
    ):
        style = document.styles[style_name]
        assert style.font.size.pt == size
        assert style.paragraph_format.space_before.pt == before
        assert style.paragraph_format.space_after.pt == after

    for table in document.tables:
        properties = table._tbl.tblPr
        assert properties.find(qn("w:tblW")).get(qn("w:w")) == "9360"
        assert properties.find(qn("w:tblInd")).get(qn("w:w")) == "120"
        grid_widths = [
            int(column.get(qn("w:w")))
            for column in table._tbl.tblGrid
        ]
        assert sum(grid_widths) == 9360
        assert len(grid_widths) == len(table.columns)
        for row in table.rows:
            assert row.height is None
            for index, cell in enumerate(row.cells):
                width = cell._tc.tcPr.find(qn("w:tcW"))
                assert int(width.get(qn("w:w"))) == grid_widths[index]


def test_docx_generation_is_hash_deterministic() -> None:
    subprocess.run(
        [sys.executable, str(GENERATOR)],
        cwd=PROJECT_ROOT,
        check=True,
        capture_output=True,
        text=True,
    )
    first = _sha256(DOCX_PATH)
    subprocess.run(
        [sys.executable, str(GENERATOR)],
        cwd=PROJECT_ROOT,
        check=True,
        capture_output=True,
        text=True,
    )
    assert _sha256(DOCX_PATH) == first
