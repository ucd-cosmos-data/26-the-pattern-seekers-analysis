from __future__ import annotations

import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from docx import Document

from src.models.outfield_tournament_impact_v4 import (
    REQUIRED_DEFENSIVE_PIPELINE_STAGES,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DIAGNOSTICS = (
    PROJECT_ROOT / "results" / "diagnostics" / "ranking_repair"
)
RANKING = PROJECT_ROOT / "results" / "reports" / "ranking"
V4_ROOT = PROJECT_ROOT / "results" / "reports" / "v4"
AUDIT_PATH = DIAGNOSTICS / "outfield_v4_release_audit.json"
PREREG_PATH = DIAGNOSTICS / "outfield_v4_preregistration.json"
RICH_PATH = RANKING / "player_rankings_v4.csv"
TOP50_PATH = DIAGNOSTICS / "outfield_v4_top50.csv"
MOVEMENT_PATH = DIAGNOSTICS / "outfield_v4_fixture_movement.csv"
GRID_PATH = DIAGNOSTICS / "outfield_v4_configuration_grid.csv"
DOCX_PATH = V4_ROOT / "docs" / "final_summary_v4.docx"
DOCX_VALIDATION_PATH = DIAGNOSTICS / "outfield_v4_docx_validation.json"
V4_HASH_PATH = DIAGNOSTICS / "outfield_v4_artifact_hashes.json"
MASTER_MANIFEST_PATH = (
    PROJECT_ROOT / "results" / "metadata" / "artifact_manifest.json"
)
ACTIVE_MODEL_VERSION = (
    "outfield-tournament-impact-v4+goalkeeper-event-profile-v3"
)
pytestmark = pytest.mark.skipif(
    not V4_ROOT.exists(),
    reason="Superseded v4 report tree was intentionally removed",
)
REMOVED_REDUNDANT_RANKING_ALIASES = {
    "results/reports/ranking/global_rankings_outfield_300min.csv",
    "results/reports/ranking/goalkeeper_rankings_unified.csv",
    "results/reports/ranking/outfield_rankings_v4.csv",
    "results/reports/ranking/outfield_rankings_v4.json",
    "results/reports/ranking/player_rankings_v2.csv",
    "results/reports/ranking/v5_player_rankings.csv",
    "results/reports/ranking/v5_player_rankings.json",
}


def _json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def test_preregistration_is_frozen_and_complete_before_results() -> None:
    prereg = _json(PREREG_PATH)
    assert prereg["written_before_final_ranking"] is True
    assert prereg["written_before_any_v4_results"] is True
    assert prereg["contains_final_ranking_results"] is False
    assert prereg["configuration_selection_status"] == "not_started"
    assert tuple(prereg["defensive_pipeline_stages"]) == (
        REQUIRED_DEFENSIVE_PIPELINE_STAGES
    )
    grid = prereg["candidate_grid"]
    assert grid["variance_share_target"] == [0.35, 0.42, 0.5]
    candidate_count = int(
        np.prod([len(values) for values in grid.values()])
    )
    assert candidate_count == 1296


def test_release_audit_records_passed_coupled_pipeline_and_firewall() -> None:
    audit = _json(AUDIT_PATH)
    assert audit["release_status"] == "passed"
    assert tuple(audit["defensive_pipeline_stages"]) == (
        REQUIRED_DEFENSIVE_PIPELINE_STAGES
    )
    assert audit["configuration_grid"]["candidate_count"] == 1296
    assert audit["configuration_grid"]["eligible_count"] == 864
    assert audit["selected_prevention_validation"]["gate_passed"] is True
    assert (
        audit["selected_prevention_validation"]["oof_rmse"]
        < audit["selected_prevention_validation"]["null_rmse"]
    )
    firewall = audit["selection_firewall"]
    assert firewall[
        "configuration_locked_before_named_fixture_review"
    ] is True
    assert firewall["named_players_used_as_scoring_inputs"] is False
    assert firewall["named_players_used_for_grid_selection"] is False
    assert firewall["external_consensus_used_for_grid_selection"] is False
    assert firewall["advancement_or_round_bonus"] is False
    assert audit["full_test_suite"]["status"] == "passed"
    assert audit["full_test_suite"]["passed"] >= 161
    assert audit["docx_validation"]["render_validation_status"] == "passed"
    assert (
        "scripts/update_outfield_v4_final_summary_docx.py"
        in audit["source_hashes"]
    )
    limitations = "\n".join(audit["limitations"])
    assert "113.7" in limitations
    assert "maximum is 262" in limitations


def test_complete_grid_and_selected_variance_contract() -> None:
    grid = pd.read_csv(GRID_PATH)
    assert len(grid) == 1296
    assert grid["config_id"].is_unique
    assert int(grid["eligible_for_release"].sum()) == 864
    selected = grid.loc[grid["selection_status"].eq("selected")]
    assert len(selected) == 1
    row = selected.iloc[0]
    assert row["prevention_weight"] > 0.0
    assert 0.35 - 1e-12 <= row[
        "realized_defensive_variance_share"
    ] <= 0.50 + 1e-12
    assert row["bootstrap_top50_jaccard_median"] > 0.0
    assert row["leave_one_match_out_median_top50_jaccard"] > 0.0


def test_release_cohorts_ranks_scale_and_model_version() -> None:
    rich = pd.read_csv(RICH_PATH, low_memory=False)
    outfield = rich["position_group"].ne("Goalkeeper")
    main_goalkeeper = (
        rich["position_group"].eq("Goalkeeper")
        & rich["is_main_goalkeeper"].fillna(False).astype(bool)
    )
    unified = outfield | main_goalkeeper
    minutes = pd.to_numeric(rich["minutes_played"], errors="raise")
    assert len(rich) == 593
    assert int(outfield.sum()) == 553
    assert int(unified.sum()) == 585
    assert int((outfield & minutes.ge(300.0)).sum()) == 126
    assert int((unified & minutes.ge(300.0)).sum()) == 142
    assert sorted(
        rich.loc[outfield, "tournament_impact_rank_outfield_v4"]
        .astype(int)
        .tolist()
    ) == list(range(1, 554))
    assert sorted(
        rich.loc[unified, "publication_global_rank_outfield_v4"]
        .astype(int)
        .tolist()
    ) == list(range(1, 586))
    display = pd.to_numeric(
        rich.loc[unified, "Tournament Performance Score"], errors="raise"
    )
    assert display.between(55.0, 99.0).all()
    assert np.allclose(display * 10.0, np.round(display * 10.0))
    assert rich["active_model_version"].nunique() == 1
    assert "v4" in str(rich["active_model_version"].iloc[0]).lower()


def test_required_v4_fields_and_mixture_contract() -> None:
    rich = pd.read_csv(RICH_PATH, low_memory=False)
    outfield = rich.loc[rich["position_group"].ne("Goalkeeper")]
    required = {
        "tournament_impact_score_outfield_v4",
        "tournament_impact_rank_outfield_v4",
        "defensive_value_raw_v4",
        "opponent_attack_strength_faced_v4",
        "defensive_value_opposition_adjusted_v4",
        "off_ball_prevention_value_v4",
        "defensive_value_reliability_shrunk_v4",
        "defensive_value_variance_rescaled_v4",
        "attack_reliability_v4",
        "defensive_reliability_v4",
        "attacking_channel_weight_v4",
        "defending_channel_weight_v4",
        "within_position_rank_outfield_v4",
        "within_subrole_rank_outfield_v4",
        "tournament_impact_interval_low_outfield_v4",
        "tournament_impact_interval_high_outfield_v4",
    }
    assert required.issubset(rich.columns)
    attack_weight = pd.to_numeric(
        outfield["attacking_channel_weight_v4"], errors="raise"
    )
    defense_weight = pd.to_numeric(
        outfield["defending_channel_weight_v4"], errors="raise"
    )
    assert np.allclose(attack_weight + defense_weight, 1.0, atol=1e-12)
    assert attack_weight.between(0.25, 0.75).all()
    assert attack_weight.nunique() > 50


def test_every_acceptance_gate_and_named_fixture_passes() -> None:
    audit = _json(AUDIT_PATH)
    gates = audit["acceptance_gates"]
    failed = {
        name: record
        for name, record in gates.items()
        if not bool(record["passed"])
    }
    assert not failed
    movement = pd.read_csv(MOVEMENT_PATH).set_index("fixture")
    for fixture in ("van_dijk", "romero", "otamendi"):
        assert int(movement.loc[fixture, "rank_improvement"]) >= 100
    assert int(movement.loc["bellingham", "v4_publication_rank"]) <= 50
    assert int(movement.loc["messi", "v4_publication_rank"]) == 1
    assert int(movement.loc["amrabat", "rank_improvement"]) >= 100
    assert float(movement.loc["amrabat", "off_ball_prevention_v4"]) > 0.0
    assert gates["coupling_regression"]["detail"][
        "released_coupled_rank"
    ] < gates["coupling_regression"]["detail"]["v3_outfield_rank"]
    assert gates["coupling_regression"]["detail"][
        "uncoupled_raw_rescaled_rank"
    ] >= gates["coupling_regression"]["detail"]["v3_outfield_rank"]


def test_hakimi_is_moroccos_top_outfielder_and_top50_is_exact() -> None:
    rich = pd.read_csv(RICH_PATH, low_memory=False)
    morocco = rich.loc[
        rich["team"].eq("Morocco")
        & rich["position_group"].ne("Goalkeeper")
    ].sort_values("publication_global_rank_outfield_v4")
    assert morocco.iloc[0]["player_name"] == "Achraf Hakimi Mouh"
    top50 = pd.read_csv(TOP50_PATH)
    assert len(top50) == 50
    assert top50["tournament_impact_rank_outfield_v4"].astype(int).tolist() == (
        list(range(1, 51))
    )


def test_v3_preservation_and_no_identity_inputs_are_audited() -> None:
    audit = _json(AUDIT_PATH)
    gates = audit["acceptance_gates"]
    assert gates["v3_columns_preserved"]["passed"] is True
    assert gates["v3_artifacts_preserved_before_release"]["passed"] is True
    assert gates["v3_artifacts_preserved_after_release"]["passed"] is True
    assert gates["identity_shuffle_invariance"]["passed"] is True
    assert gates["no_advancement_or_named_inputs"]["passed"] is True
    assert gates["no_advancement_or_named_inputs"]["detail"][
        "prohibited_inputs_used"
    ] == []


def test_v4_owned_manifest_hashes_every_listed_artifact() -> None:
    manifest_path = V4_ROOT / "artifact_manifest.json"
    manifest = _json(manifest_path)
    assert manifest["artifact_count"] == len(manifest["artifacts"])
    assert manifest["artifact_count"] > 2000
    assert manifest["active_model_version"] == ACTIVE_MODEL_VERSION
    artifact_paths = {record["path"] for record in manifest["artifacts"]}
    assert DOCX_PATH.relative_to(PROJECT_ROOT).as_posix() in artifact_paths
    for record in manifest["artifacts"]:
        path = PROJECT_ROOT / record["path"]
        if record["path"] in REMOVED_REDUNDANT_RANKING_ALIASES:
            assert not path.exists()
            continue
        assert path.is_file()


def test_canonical_docx_is_structurally_and_visually_validated() -> None:
    validation = _json(DOCX_VALIDATION_PATH)
    assert DOCX_PATH.is_file()
    assert validation["source_release_status"] == "passed"
    assert validation["structural_validation_passed"] is True
    assert validation["render_validation_status"] == "passed"
    assert validation["rendered_pages"] == 8
    assert validation["visual_inspection"]["all_pages_inspected"] is True
    assert validation["visual_inspection"]["clipping_or_overflow"] is False
    assert validation["sha256"] == _sha256(DOCX_PATH)

    document = Document(DOCX_PATH)
    paragraph_text = "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )
    assert "Tournament Impact v4" in paragraph_text
    assert "External consensus comparison" in paragraph_text
    assert "Remaining limitations" in paragraph_text
    assert len(document.tables) == validation["tables"]
    assert len(document.inline_shapes) == validation["inline_shapes"]


def test_final_v4_hash_manifest_covers_docx_validation_and_audit() -> None:
    manifest = _json(V4_HASH_PATH)
    assert manifest["active_model_version"] == ACTIVE_MODEL_VERSION
    assert manifest["artifact_count"] == len(manifest["artifacts"])
    assert manifest["artifact_count"] > 2046
    paths = {record["path"] for record in manifest["artifacts"]}
    required = {
        DOCX_PATH.relative_to(PROJECT_ROOT).as_posix(),
        DOCX_VALIDATION_PATH.relative_to(PROJECT_ROOT).as_posix(),
        AUDIT_PATH.relative_to(PROJECT_ROOT).as_posix(),
        (V4_ROOT / "artifact_manifest.json").relative_to(
            PROJECT_ROOT
        ).as_posix(),
    }
    assert required.issubset(paths)
    for record in manifest["artifacts"]:
        path = PROJECT_ROOT / record["path"]
        if record["path"] in REMOVED_REDUNDANT_RANKING_ALIASES:
            assert not path.exists()
            continue
        assert path.is_file()


def test_global_catalog_tracks_promoted_release_and_hashes_docx() -> None:
    manifest = _json(MASTER_MANIFEST_PATH)
    assert manifest["active_model_version"] == (
        "ranking-repair-v3.0-qatar-2022"
    )
    docx_record = next(
        record
        for record in manifest["artifacts"]
        if record["path"] == DOCX_PATH.relative_to(
            PROJECT_ROOT
        ).as_posix()
    )
    assert docx_record["sha256"] == _sha256(DOCX_PATH)
